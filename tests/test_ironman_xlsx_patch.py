"""Offline unit regressions; no commercial workbook or TrainingPeaks credentials needed."""
from __future__ import annotations
import hashlib
import sys
from pathlib import Path
from tempfile import TemporaryDirectory
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'scripts' / 'ironman_jobs'))
# Local standalone test execution from the work directory also works.
sys.path.insert(0, str(Path(__file__).resolve().parent))
from ironman_xlsx_patch import SCHEMA, PatchStop, apply, prepare, discover

@pytest.fixture
def book(tmp_path):
    p=tmp_path/'synthetic.xlsx'
    w=Workbook();a=w.active;a.title='GIORNI 168'
    a['A1']='W16';a['B1']='GIO';a['C1']='HI RICHIAMO 6×1′';a['D1']=140
    a['E1']='=D1*2';a['F1']=1
    a['A2']='W17';a['C2']='TEST';a['D2']=150
    b=w.create_sheet('BICI QUALITA 24W');b['A1']='W16';b['B1']='HI RICHIAMO 6×1′'
    chart=BarChart();chart.add_data(Reference(a, min_col=4,max_col=4,min_row=1,max_row=2));a.add_chart(chart,'H1')
    w.save(p)
    return p

def manifest(p, changes, assertions=(),**extras):
    return dict(schema=SCHEMA,input_sha256=hashlib.sha256(p.read_bytes()).hexdigest(),
                scope={'resource':'xlsx','week':'W16'},changes=changes,assertions=list(assertions),**extras)

def item(sheet,cell,before,after):return dict(sheet=sheet,cell=cell,before=before,after=after)

def test_one_cell_writes_only_local_part(book,tmp_path):
    m=manifest(book,[item('GIORNI 168','C1','HI RICHIAMO 6×1′','HI RICHIAMO 6×1′ · selettivo')])
    r=apply(book,m,tmp_path/'edited.xlsx')
    assert r['route']=='A_LOCAL' and r['cells']==1 and r['xlsx_parts_changed_count']==1
    assert r['formula_caches_preserved']==1 and r['chart_parts_preserved']==1
    assert load_workbook(tmp_path/'edited.xlsx')['GIORNI 168']['C1'].value.endswith('selettivo')

def test_mirrors_and_coverage(book,tmp_path):
    m=manifest(book,[item('GIORNI 168','C1','HI RICHIAMO 6×1′','HI ESPOSIZIONE 6×1′'),
                     item('BICI QUALITA 24W','B1','HI RICHIAMO 6×1′','HI ESPOSIZIONE 6×1′')],
               [{'sheet':'GIORNI 168','week':'W16','contains':'RICHIAMO','before_count':1,'after_count':0},
                {'sheet':'BICI QUALITA 24W','week':'W16','contains':'RICHIAMO','before_count':1,'after_count':0}])
    r=apply(book,m,tmp_path/'mirrors.xlsx');assert r['cells']==2 and r['sheets_modified']==2

def test_dose_change_blocked(book,tmp_path):
    m=manifest(book,[item('GIORNI 168','C1','HI RICHIAMO 6×1′','HI RICHIAMO 8×3′')])
    assert prepare(book,m)['route']=='B_BLOCK_REVIEW_STOP'
    with pytest.raises(PatchStop,match='ESCALATION_REQUIRED'):apply(book,m,tmp_path/'not_created.xlsx')
    assert not (tmp_path/'not_created.xlsx').exists()

def test_intensity_change_without_numbers_blocked(book,tmp_path):
    m=manifest(book,[item('GIORNI 168','C1','HI RICHIAMO 6×1′','EASY RICHIAMO 6×1′')])
    assert prepare(book,m)['route']=='B_BLOCK_REVIEW_STOP'
    with pytest.raises(PatchStop,match='ESCALATION_REQUIRED'):apply(book,m,tmp_path/'not_created.xlsx')

def test_numeric_driver_blocked(book,tmp_path):
    m=manifest(book,[item('GIORNI 168','D1',140,160)])
    with pytest.raises(PatchStop,match='TEXT_ONLY_V1'):apply(book,m,tmp_path/'not_created.xlsx')

def test_version_and_wrong_week_guard(book,tmp_path):
    m=manifest(book,[item('GIORNI 168','C2','TEST','TEST nuovo')])
    with pytest.raises(PatchStop,match='WEEK_SCOPE_MISMATCH'):apply(book,m,tmp_path/'badweek.xlsx')
    m=manifest(book,[item('GIORNI 168','C1','HI RICHIAMO 6×1′','note')]);m['input_sha256']='0'*64
    with pytest.raises(PatchStop,match='SOURCE_HASH_DRIFT'):apply(book,m,tmp_path/'badsha.xlsx')

def test_scope_across_many_weeks_escalates(book):
    m=manifest(book,[item('GIORNI 168','C1','HI RICHIAMO 6×1′','HI ESPOSIZIONE 6×1′'),
                     item('GIORNI 168','C2','TEST','ESAME')])
    m['scope'].pop('week')
    m['changes'][0]['week']='W16';m['changes'][1]['week']='W17'
    assert prepare(book,m)['route']=='B_BLOCK_REVIEW_STOP'

def test_inspection_new_context(book):
    d=discover(book,'W16','RICHIAMO')
    assert d['matches']==2 and d['sheets_consulted']==2
