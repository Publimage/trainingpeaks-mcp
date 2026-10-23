#!/usr/bin/env python3
"""IRONMAN workbook surgical patcher, V1. Offline, fail-closed, zero TrainingPeaks access.

CLI: inspect, preview, apply. A change is an exact cell preimage and replacement,
not a free-text, whole-workbook replace. ZIP members outside altered worksheets
are preserved byte-for-byte (uncompressed content); formula caches and charts
must be unchanged. A separate coach review is required for risk escalations.

The Google Drive KB is the only canonical methodological source: this is just a
technical executor for reviewed changes to disposable/specified XLSX copies.
"""
from __future__ import annotations

import argparse
import copy
import hashlib
import json
import os
import re
import tempfile
import time
import unicodedata
import zipfile
from collections import Counter
from pathlib import Path
from xml.etree import ElementTree as ET

from openpyxl import load_workbook

SCHEMA = "IRONMAN_XLSX_PATCH_V1"
NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PKG = "http://schemas.openxmlformats.org/package/2006/relationships"
ET.register_namespace("", NS)
ET.register_namespace("r", REL)
CELL = re.compile(r"^[A-Z]{1,3}[1-9][0-9]{0,6}$")
NUMBER = re.compile(r"\d+(?:[.,]\d+)?")
WEEK = re.compile(r"^W(\d{1,2})$", re.IGNORECASE)
DANGER = re.compile(r"(?i)(taper|periodizz|macrocicl|race.week|soglia|threshold|ftp|vo.?2|hi-|maxim|test|retest|recuper|durat|minuti|volume|intensit|caric|\bcss\b)")

class PatchStop(RuntimeError):
    pass

def require(condition, message):
    if not condition:
        raise PatchStop(message)

def digest(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()

def sheets_to_xml(z):
    wb = ET.fromstring(z.read("xl/workbook.xml"))
    rels = ET.fromstring(z.read("xl/_rels/workbook.xml.rels"))
    targets = {r.attrib["Id"]: r.attrib["Target"] for r in rels}
    mapping = {}
    for sh in wb.find("{%s}sheets" % NS):
        target = targets[sh.attrib["{%s}id" % REL]]
        if target.startswith("/"):
            part = target.lstrip("/")
        else:
            from posixpath import normpath, join
            part = normpath(join("xl", target))
        require(part in z.namelist(), f"XML_PART_MISSING:{sh.attrib['name']}")
        mapping[sh.attrib["name"]] = part
    return mapping

def all_formula_cache_fingerprint(z):
    fingerprints = []
    for sname, part in sheets_to_xml(z).items():
        root = ET.fromstring(z.read(part))
        for c in root.iter("{%s}c" % NS):
            f = c.find("{%s}f" % NS)
            if f is not None:
                v = c.find("{%s}v" % NS)
                fingerprints.append((sname, c.attrib.get("r"), f.attrib, f.text, None if v is None else v.text))
    return fingerprints

def load_input(path):
    require(Path(path).is_file(), "INPUT_NOT_FOUND")
    raw = Path(path).read_bytes()
    with zipfile.ZipFile(path) as z:
        require(z.testzip() is None, "BROKEN_ZIP")
    return raw, load_workbook(path, data_only=False, read_only=False)

def week_rows(ws, week):
    if not week:
        return range(1, ws.max_row + 1)
    tag = f"W{int(str(week).upper().lstrip('W')):02d}"
    return [r for r in range(1, ws.max_row+1) if str(ws.cell(r,1).value).upper().strip() == tag]

def discover(path, week, term, limit=120):
    t0 = time.monotonic()
    _, wb = load_input(path)
    out = []
    for ws in wb:
        for r in week_rows(ws, week):
            for c in ws[r]:
                v = c.value
                if isinstance(v, str) and c.data_type != "f" and (not term or term.casefold() in v.casefold()):
                    out.append({"sheet":ws.title, "cell":c.coordinate, "value":v[:320]})
    return {"schema":SCHEMA,"mode":"inspect","week":week,"term":term,
            "matches":len(out),"items":out[:limit],"truncated":len(out)>limit,
            "sheets_consulted":len(wb.worksheets),"seconds":round(time.monotonic()-t0,4)}

def normalized(value):
    return unicodedata.normalize("NFKC", str(value)).casefold()


def numeric_changed(before, after):
    return NUMBER.findall(normalized(before)) != NUMBER.findall(normalized(after))


STIMULUS = re.compile(
    r"\b(?:easy|recovery|recupero|medio|tempo|soglia|threshold|ftp|vo2|max|hi|lt1|lt2|"
    r"rpe|test|retest|taper|race|fartlek|progressivo|forza|strength|scarico)\b"
)


def stimulus_signature(value):
    return Counter(STIMULUS.findall(normalized(value)))


def classify(wb, changes, manifest):
    weeks = {str(x.get("week",manifest.get("scope",{}).get("week", ""))).upper() for x in changes}
    weeks.discard("")
    reasons = []
    if len(weeks) > 1:
        reasons.append("MULTI_WEEK_SCOPE")
    for ch in changes:
        a,b = ch["before"],ch["after"]
        if stimulus_signature(a)!=stimulus_signature(b):
            reasons.append(f"STIMULUS_SEMANTICS_CHANGED:{ch['sheet']}!{ch['cell']}")
        if numeric_changed(a,b): reasons.append(f"NUMERIC_DOSE_OR_VOLUME:{ch['sheet']}!{ch['cell']}")
        if isinstance(a,(int,float)) or isinstance(b,(int,float)):
            reasons.append(f"NON_TEXT_VALUE:{ch['sheet']}!{ch['cell']}")
        if (DANGER.search(str(a)) or DANGER.search(str(b))) and a!=b and numeric_changed(a,b):
            reasons.append(f"PHYSIOLOGY_OR_TEST_DEPENDENCY:{ch['sheet']}!{ch['cell']}")
        if ch.get("impact") in ("block","structural"):
            reasons.append(f"DECLARED_{ch['impact'].upper()}:{ch['sheet']}!{ch['cell']}")
        # Formula/chart drivers cannot be replaced by a text-only executor.
        ws=wb[ch["sheet"]];cell=ws[ch["cell"]]
        if cell.data_type == "f":reasons.append(f"FORMULA_CELL:{ch['sheet']}!{ch['cell']}")
    if len(weeks) >= 12 or any("STRUCTURAL" in x for x in reasons):
        route = "C_STRUCTURAL_STOP"
    elif reasons:
        route = "B_BLOCK_REVIEW_STOP"
    else:
        route = "A_LOCAL"
    # A large number of purely textual mirrors is not itself structural.
    return route, sorted(set(reasons))

def prepare(input_path, manifest):
    t0=time.monotonic()
    raw, wb = load_input(input_path)
    require(isinstance(manifest,dict) and manifest.get("schema")==SCHEMA,"MANIFEST_SCHEMA")
    require(manifest.get("input_sha256")==digest(raw),"SOURCE_HASH_DRIFT")
    scope=manifest.get("scope")
    require(isinstance(scope,dict) and scope.get("resource")=="xlsx", "RESOURCE_SCOPE")
    changes=manifest.get("changes")
    require(isinstance(changes,list) and 1<=len(changes)<=250,"CHANGE_COUNT")
    seen=set();diffs=[]
    for ch in changes:
        require(isinstance(ch,dict),"CHANGE_NOT_OBJECT")
        require(all(k in ch for k in ("sheet","cell","before","after")),"INCOMPLETE_CHANGE")
        sn,coord=ch["sheet"],ch["cell"]
        require(sn in wb.sheetnames and CELL.fullmatch(coord),"UNKNOWN_SHEET_OR_CELL")
        require((sn,coord) not in seen,"DUPLICATE_CELL")
        seen.add((sn,coord))
        c=wb[sn][coord]
        require(c.value==ch["before"],f"PREIMAGE_MISMATCH:{sn}!{coord}: {c.value!r}")
        require(ch["before"]!=ch["after"],f"NOOP:{sn}!{coord}")
        require(c.data_type!="f",f"FORMULA_PATCH_FORBIDDEN:{sn}!{coord}")
        require(c.data_type not in ("e",),f"ERROR_CELL:{sn}!{coord}")
        require(isinstance(ch["before"],str) and isinstance(ch["after"],str), f"TEXT_ONLY_V1:{sn}!{coord}")
        require(bool(ch["after"].strip()),f"EMPTY_RESULT:{sn}!{coord}")
        if scope.get("week"):
            require(c.row in week_rows(wb[sn],scope["week"]),f"WEEK_SCOPE_MISMATCH:{sn}!{coord}")
        diffs.append({"sheet":sn,"cell":coord,"before":ch["before"],"after":ch["after"]})
    route,reasons=classify(wb,changes,manifest)
    expected_route=manifest.get("expected_route")
    if expected_route:require(expected_route==route,"ROUTE_UNEXPECTED:"+route)
    # Before assertions allow the manifest author to prove exact mirror coverage.
    for a in manifest.get("assertions",[]):
        sn=a.get("sheet"); week=a.get("week",scope.get("week")); term=a.get("contains")
        require(sn in wb.sheetnames and isinstance(term,str) and term,"BAD_ASSERTION")
        num=sum(term in str(c.value) for r in week_rows(wb[sn],week) for c in wb[sn][r] if c.value is not None)
        require(num==a.get("before_count"),f"ASSERTION_COVERAGE_DRIFT:{sn}:{term}:{num}")
    return {"input":str(input_path),"sha256_before":digest(raw),"scope":scope,"route":route,
            "reasons":reasons,"diffs":diffs,"cells":len(diffs),"sheets_modified":len({x['sheet'] for x in diffs}),
            "sheets_inspected":len(wb.worksheets),"seconds_preflight":round(time.monotonic()-t0,4)}

def patch_sheet(xml, edits):
    root=ET.fromstring(xml)
    byref={c.attrib.get("r"):c for c in root.iter("{%s}c" % NS)}
    for coord, new in edits.items():
        require(coord in byref, "MISSING_CELL_XML:"+coord)
        el=byref[coord]
        require(el.find("{%s}f" % NS) is None,"FORMULA_CELL_XML:"+coord)
        # A rich-text run can contain intentionally formatted portions; stop rather than flatten.
        ist=el.find("{%s}is" % NS)
        if ist is not None:
            require(not list(ist.findall("{%s}r" % NS)),"RICH_TEXT_NOT_SUPPORTED:"+coord)
        for node in list(el):
            if node.tag in ("{%s}is" % NS,"{%s}v" % NS): el.remove(node)
        el.attrib["t"]="inlineStr"
        inline=ET.SubElement(el,"{%s}is" % NS)
        val=ET.SubElement(inline,"{%s}t" % NS)
        if new != new.strip() or "\n" in new:val.set("{http://www.w3.org/XML/1998/namespace}space","preserve")
        val.text=new
    return ET.tostring(root,encoding="utf-8",xml_declaration=False)

def validate_after(input_path, output_path, prepared, manifest):
    before=load_workbook(input_path,data_only=False)
    after=load_workbook(output_path,data_only=False)
    require(before.sheetnames==after.sheetnames,"SHEET_SET_CHANGED")
    targets={(x["sheet"],x["cell"]):x["after"] for x in prepared["diffs"]}
    total=0
    for a,b in zip(before,after):
        require((a.max_row,a.max_column)==(b.max_row,b.max_column),"SHEET_DIMENSIONS_CHANGED:"+a.title)
        require(len(a._charts)==len(b._charts),"CHART_COUNT_CHANGED:"+a.title)
        require([str(x) for x in a.merged_cells.ranges]==[str(x) for x in b.merged_cells.ranges],"MERGES_CHANGED:"+a.title)
        for row in a:
            for c in row:
                d=b[c.coordinate]; k=(a.title,c.coordinate)
                require(d.value==(targets[k] if k in targets else c.value),"CELL_DRIFT:"+str(k))
                require(d.style_id==c.style_id,"STYLE_CHANGED:"+str(k))
                total+=1
    for a in manifest.get("assertions",[]):
        sn=a["sheet"];week=a.get("week",manifest["scope"].get("week")); term=a["contains"]
        count=sum(term in str(c.value) for r in week_rows(after[sn],week) for c in after[sn][r] if c.value is not None)
        require(count==a.get("after_count"),f"POST_ASSERTION_FAILED:{sn}:{term}:{count}")
    with zipfile.ZipFile(input_path) as iz,zipfile.ZipFile(output_path) as oz:
        require(set(iz.namelist())==set(oz.namelist()),"ZIP_ENTRIES_CHANGED")
        allowed={sheets_to_xml(iz)[x["sheet"]] for x in prepared["diffs"]}
        diffparts=[]
        for p in iz.namelist():
            if iz.read(p)!=oz.read(p):diffparts.append(p)
        require(set(diffparts)==allowed,"UNEXPECTED_ZIP_PART_DRIFT:"+str(set(diffparts)-allowed))
        before_formula=all_formula_cache_fingerprint(iz)
        after_formula=all_formula_cache_fingerprint(oz)
        require(before_formula==after_formula,"FORMULA_OR_CACHE_DRIFT")
        charts=[x for x in iz.namelist() if x.startswith("xl/charts/")]
    return {"xlsx_parts_changed":diffparts,"xlsx_parts_changed_count":len(diffparts),
            "formula_caches_preserved":len(before_formula),"chart_parts_preserved":len(charts),
            "cells_checked":total,"bytes_sha256_after":digest(Path(output_path).read_bytes()),
            "integrity":"PASS"}

def apply(input_path, manifest, output_path):
    t0=time.monotonic()
    prepared=prepare(input_path,manifest)
    require(prepared["route"]=="A_LOCAL","ESCALATION_REQUIRED:"+prepared["route"]+":"+",".join(prepared["reasons"]))
    require(Path(input_path).resolve()!=Path(output_path).resolve(),"SOURCE_OVERWRITE_FORBIDDEN")
    require(not Path(output_path).exists(),"OUTPUT_EXISTS")
    Path(output_path).parent.mkdir(parents=True,exist_ok=True)
    with zipfile.ZipFile(input_path) as iz:
        mapping=sheets_to_xml(iz)
        edits={}
        for x in prepared["diffs"]:
            edits.setdefault(mapping[x["sheet"]],{})[x["cell"]]=x["after"]
        # Write in a sibling temporary file. Never leave a partially-written result.
        fd,tmp=tempfile.mkstemp(suffix=".xlsx",prefix=".ironman-patch-",dir=str(Path(output_path).parent))
        os.close(fd)
        try:
            with zipfile.ZipFile(tmp,"w") as oz:
                oz.comment=iz.comment
                for info in iz.infolist():
                    payload=iz.read(info.filename)
                    if info.filename in edits:payload=patch_sheet(payload,edits[info.filename])
                    oz.writestr(copy.copy(info),payload)
            verified=validate_after(input_path,tmp,prepared,manifest)
            os.replace(tmp,output_path)
        finally:
            if Path(tmp).exists():os.unlink(tmp)
    prepared.update(verified)
    prepared["output"]=str(output_path)
    prepared["mode"]="apply"
    prepared["seconds_total"]=round(time.monotonic()-t0,4)
    prepared["provider_writes"]=0
    prepared["google_drive_writes"]=0
    return prepared

def main():
    ap=argparse.ArgumentParser(description=__doc__)
    sub=ap.add_subparsers(dest="command",required=True)
    i=sub.add_parser("inspect");i.add_argument("--input",required=True);i.add_argument("--week");i.add_argument("--term",default="");i.add_argument("--limit",type=int,default=120)
    for name in ("preview","apply"):
        p=sub.add_parser(name);p.add_argument("--input",required=True);p.add_argument("--manifest",required=True)
        if name=="apply":p.add_argument("--output",required=True)
    ap.add_argument("--report",help="save machine-readable JSON report")
    ns=ap.parse_args()
    try:
        if ns.command=="inspect":result=discover(ns.input,ns.week,ns.term,ns.limit)
        else:
            m=json.loads(Path(ns.manifest).read_text(encoding="utf-8"))
            result=prepare(ns.input,m) if ns.command=="preview" else apply(ns.input,m,ns.output)
            result["mode"]=ns.command
            result["provider_writes"]=0
        rc=0 if result.get("route") not in ("B_BLOCK_REVIEW_STOP","C_STRUCTURAL_STOP") else 2
    except (PatchStop,ValueError,KeyError,zipfile.BadZipFile) as exc:
        result={"status":"STOP","error":str(exc),"provider_writes":0,"mode":ns.command}
        rc=2
    if ns.report:Path(ns.report).write_text(json.dumps(result,ensure_ascii=False,indent=2,default=str),encoding="utf-8")
    print(json.dumps(result,ensure_ascii=False,indent=2,default=str))
    raise SystemExit(rc)

if __name__=="__main__":main()
